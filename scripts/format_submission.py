#!/usr/bin/env python3
"""Apply the final EarnScore submission layout to a Pandoc-generated DOCX."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "EAF0F7"
VERY_LIGHT = "F5F7FA"
MID_GREY = "667085"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row) -> None:
    repeat_header(row)


def remove_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def set_run_font(run, name="Aptos", size=None, color=None, bold=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def style_document(source: Path, destination: Path) -> None:
    doc = Document(source)

    properties = doc.core_properties
    properties.title = "EarnScore — Alternative Creditworthiness Score for Gig Workers"
    properties.subject = "FinTech hackathon submission"
    properties.author = "Winterfell Club"
    properties.keywords = "EarnScore, gig workers, alternative credit scoring, financial inclusion"

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("202939")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Title": (28, NAVY, 18, 12),
        "Subtitle": (16, BLUE, 5, 10),
        "Heading 1": (16, NAVY, 14, 7),
        "Heading 2": (13, BLUE, 11, 5),
        "Heading 3": (11.5, NAVY, 9, 4),
        "Heading 4": (10.5, BLUE, 7, 3),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name in {"Title", "Subtitle", "Heading 1"} else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Block Text" in doc.styles:
        block = doc.styles["Block Text"]
        block.font.name = "Aptos"
        block.font.size = Pt(10.5)
        block.font.color.rgb = RGBColor.from_string(NAVY)
        block.paragraph_format.left_indent = Inches(0.25)
        block.paragraph_format.right_indent = Inches(0.25)
        block.paragraph_format.space_before = Pt(8)
        block.paragraph_format.space_after = Pt(8)

    if "Source Code" in doc.styles:
        code = doc.styles["Source Code"]
        code.font.name = "Aptos Mono"
        code.font.size = Pt(9.5)
        code.font.color.rgb = RGBColor.from_string(NAVY)
        code.paragraph_format.left_indent = Inches(0.35)
        code.paragraph_format.right_indent = Inches(0.35)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(4)
        code.paragraph_format.keep_together = True

    # Cover-page treatment: all content before the first numbered section.
    first_section_index = None
    team_names = {"Nikhil Upadhyay", "Sharath B S", "Shaikh Asgar"}
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text == "1. Problem Statement":
            first_section_index = index
            paragraph.paragraph_format.page_break_before = True
            break

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.right_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(8)

        if text == "EarnScore":
            paragraph.paragraph_format.space_before = Pt(62)
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                set_run_font(run, "Aptos Display", 30, NAVY, True)
        elif text == "Alternative Creditworthiness Score for Gig Workers":
            paragraph.paragraph_format.space_after = Pt(28)
            for run in paragraph.runs:
                set_run_font(run, "Aptos Display", 17, BLUE, True)
        elif text in team_names:
            remove_numbering(paragraph)
            paragraph.style = doc.styles["Normal"]
            for run in paragraph.runs:
                set_run_font(run, "Aptos", 11, NAVY, False)
        elif text.startswith("Team name:"):
            paragraph.paragraph_format.space_before = Pt(18)
        elif text.startswith("Prototype scope:"):
            paragraph.paragraph_format.space_before = Pt(18)
            for run in paragraph.runs:
                set_run_font(run, "Aptos", 9.5, MID_GREY, run.bold)
        elif "One-line pitch:" in text:
            paragraph.paragraph_format.space_before = Pt(12)
            for run in paragraph.runs:
                set_run_font(run, "Aptos", 10.5, NAVY, run.bold)

    # Body paragraph consistency and heading safeguards.
    start = first_section_index or 0
    for paragraph in doc.paragraphs[start:]:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        elif style_name not in {"Source Code", "Caption"}:
            paragraph.paragraph_format.widow_control = True

        if "CSV / simulated gig records" in paragraph.text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Inches(0.55)
            paragraph.paragraph_format.right_indent = Inches(0.55)
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(10)
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_BLUE)
            p_pr.append(shd)

    # Professional tables with repeating headers and non-splitting rows.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if "Table Grid" in doc.styles:
            table.style = "Table Grid"
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                set_cell_shading(cell, NAVY if row_index == 0 else (LIGHT_BLUE if row_index % 2 == 0 else WHITE))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        set_run_font(run, "Aptos", 9.2, WHITE if row_index == 0 else "202939", row_index == 0 or run.bold)

    # Header and footer.
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = "EARNSCORE  |  WINTERFELL CLUB"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_run_font(run, "Aptos", 8, MID_GREY, True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    prefix = fp.add_run("EarnScore  •  ")
    set_run_font(prefix, "Aptos", 8, MID_GREY, False)
    add_field(fp, "PAGE")
    suffix = fp.add_run("  •  Hackathon Submission")
    set_run_font(suffix, "Aptos", 8, MID_GREY, False)

    # Keep field results and links current when Word/LibreOffice opens the file.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: format_submission.py INPUT.docx OUTPUT.docx")
    style_document(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
